import io
import openpyxl
from app.tasks.celery_app import celery_app
from app.services.s3_service import download_file_to_memory
from app.database import supabase_client

@celery_app.task(name="parse_excel_metadata")
def parse_excel_metadata(file_id: str, s3_key: str, user_id: str):
    try:
        # Download file bytes
        file_bytes = download_file_to_memory(s3_key)
        
        # Parse Excel file
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        
        sheet_info = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sheet_info[sheet_name] = {
                "max_row": sheet.max_row,
                "max_column": sheet.max_column
            }
            
        metadata = {
            "sheet_names": wb.sheetnames,
            "dimensions": sheet_info
        }
        
        # Update file record
        supabase_client.table("files").update({"metadata": metadata}).eq("id", file_id).execute()
        
        return {"status": "success", "file_id": file_id}
        
    except Exception as e:
        # In a real app we might update the file status to 'error'
        return {"status": "error", "error": str(e)}

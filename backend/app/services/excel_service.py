import io
import csv
import openpyxl
from openpyxl.utils import get_column_letter
from app.services.s3_service import download_file_to_memory, upload_bytes_to_s3

def _load_workbook(s3_key: str, data_only: bool = False):
    file_bytes = download_file_to_memory(s3_key)
    return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=data_only)

def _save_and_upload(wb: openpyxl.Workbook, s3_key: str):
    output = io.BytesIO()
    wb.save(output)
    upload_bytes_to_s3(output.getvalue(), s3_key)

def _execute_operation(s3_key, operation_func, requires_save=True, **kwargs):
    try:
        wb = _load_workbook(s3_key)
        result = operation_func(wb, **kwargs)
        if requires_save:
            _save_and_upload(wb, s3_key)
        return {"success": True, "result": result, "error": None}
    except Exception as e:
        return {"success": False, "result": None, "error": str(e)}

def get_workbook_info(s3_key: str):
    def op(wb):
        info = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            info[sheet_name] = {"max_row": sheet.max_row, "max_column": sheet.max_column}
        return {"sheets": wb.sheetnames, "dimensions": info}
    return _execute_operation(s3_key, op, requires_save=False)

def export_to_csv(s3_key: str, sheet_name: str):
    def op(wb):
        sheet = wb[sheet_name]
        output = io.StringIO()
        writer = csv.writer(output)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(row)
        return output.getvalue()
    # For export to CSV we want data_only=True usually, but sticking to standard structure
    try:
        wb = _load_workbook(s3_key, data_only=True)
        result = op(wb)
        return {"success": True, "result": result, "error": None}
    except Exception as e:
        return {"success": False, "result": None, "error": str(e)}

def export_to_new_xlsx(s3_key: str):
    try:
        file_bytes = download_file_to_memory(s3_key)
        return {"success": True, "result": file_bytes, "error": None}
    except Exception as e:
        return {"success": False, "result": None, "error": str(e)}

def list_sheets(s3_key: str):
    return _execute_operation(s3_key, lambda wb: wb.sheetnames, requires_save=False)

def add_sheet(s3_key: str, sheet_name: str):
    def op(wb):
        wb.create_sheet(title=sheet_name)
        return True
    return _execute_operation(s3_key, op)

def delete_sheet(s3_key: str, sheet_name: str):
    def op(wb):
        del wb[sheet_name]
        return True
    return _execute_operation(s3_key, op)

def rename_sheet(s3_key: str, old_name: str, new_name: str):
    def op(wb):
        wb[old_name].title = new_name
        return True
    return _execute_operation(s3_key, op)

def copy_sheet(s3_key: str, sheet_name: str, new_name: str):
    def op(wb):
        source = wb[sheet_name]
        target = wb.copy_worksheet(source)
        target.title = new_name
        return True
    return _execute_operation(s3_key, op)

def read_range(s3_key: str, sheet_name: str, cell_range: str):
    def op(wb):
        sheet = wb[sheet_name]
        data = []
        for row in sheet[cell_range]:
            if not isinstance(row, tuple):
                row = (row,)
            data.append([cell.value for cell in row])
        return data
    try:
        wb = _load_workbook(s3_key, data_only=True)
        result = op(wb)
        return {"success": True, "result": result, "error": None}
    except Exception as e:
        return {"success": False, "result": None, "error": str(e)}

def write_range(s3_key: str, sheet_name: str, cell_range: str, values: list):
    def op(wb):
        sheet = wb[sheet_name]
        cells = sheet[cell_range]
        if not isinstance(cells, tuple):
            cells = (cells,)
        for r_idx, row in enumerate(cells):
            if r_idx < len(values):
                if not isinstance(row, tuple):
                    row = (row,)
                for c_idx, cell in enumerate(row):
                    if c_idx < len(values[r_idx]):
                        cell.value = values[r_idx][c_idx]
        return True
    return _execute_operation(s3_key, op)

def write_cell(s3_key: str, sheet_name: str, cell: str, value):
    def op(wb):
        wb[sheet_name][cell].value = value
        return True
    return _execute_operation(s3_key, op)

def clear_range(s3_key: str, sheet_name: str, cell_range: str):
    def op(wb):
        sheet = wb[sheet_name]
        for row in sheet[cell_range]:
            if not isinstance(row, tuple):
                row = (row,)
            for cell in row:
                cell.value = None
        return True
    return _execute_operation(s3_key, op)

def merge_cells(s3_key: str, sheet_name: str, cell_range: str):
    def op(wb):
        wb[sheet_name].merge_cells(cell_range)
        return True
    return _execute_operation(s3_key, op)

def unmerge_cells(s3_key: str, sheet_name: str, cell_range: str):
    def op(wb):
        wb[sheet_name].unmerge_cells(cell_range)
        return True
    return _execute_operation(s3_key, op)

def format_cells(s3_key: str, sheet_name: str, cell_range: str, format_options: dict):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    def op(wb):
        sheet = wb[sheet_name]
        cells = sheet[cell_range]
        if not isinstance(cells, tuple):
            cells = (cells,)
            
        font_kwargs = {}
        if "bold" in format_options: font_kwargs["bold"] = format_options["bold"]
        if "italic" in format_options: font_kwargs["italic"] = format_options["italic"]
        if "font_color" in format_options: font_kwargs["color"] = format_options["font_color"].replace("#", "")
        if "font_size" in format_options: font_kwargs["size"] = format_options["font_size"]
        if "font_name" in format_options: font_kwargs["name"] = format_options["font_name"]
        
        fill = None
        if "bg_color" in format_options:
            fill = PatternFill(start_color=format_options["bg_color"].replace("#", ""), fill_type="solid")
            
        alignment = None
        if "alignment" in format_options:
            alignment = Alignment(horizontal=format_options["alignment"])
            
        font = Font(**font_kwargs) if font_kwargs else None
        
        for row in cells:
            if not isinstance(row, tuple):
                row = (row,)
            for cell in row:
                if font: cell.font = font
                if fill: cell.fill = fill
                if alignment: cell.alignment = alignment
                if "number_format" in format_options: cell.number_format = format_options["number_format"]
        return True
    return _execute_operation(s3_key, op)

def apply_formula(s3_key: str, sheet_name: str, cell: str, formula: str):
    return write_cell(s3_key, sheet_name, cell, formula)

def get_cell_formula(s3_key: str, sheet_name: str, cell: str):
    def op(wb):
        val = wb[sheet_name][cell].value
        if isinstance(val, str) and val.startswith("="):
            return val
        return None
    return _execute_operation(s3_key, op, requires_save=False)

def sort_range(s3_key: str, sheet_name: str, cell_range: str, sort_column: int, ascending: bool):
    # Implementing sorting via python logic is complex (openpyxl sort state is metadata)
    # We will sort the data directly and write it back
    def op(wb):
        sheet = wb[sheet_name]
        data = []
        cells = sheet[cell_range]
        for row in cells:
            data.append([cell.value for cell in row])
        data.sort(key=lambda x: x[sort_column] if x[sort_column] is not None else "", reverse=not ascending)
        for r_idx, row in enumerate(cells):
            for c_idx, cell in enumerate(row):
                cell.value = data[r_idx][c_idx]
        return True
    return _execute_operation(s3_key, op)

def find_and_replace(s3_key: str, sheet_name: str, find_value: str, replace_value: str):
    def op(wb):
        sheet = wb[sheet_name]
        count = 0
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and find_value in cell.value:
                    cell.value = cell.value.replace(find_value, replace_value)
                    count += 1
        return count
    return _execute_operation(s3_key, op)

def add_data_validation(s3_key: str, sheet_name: str, cell_range: str, validation_type: str, options: list):
    from openpyxl.worksheet.datavalidation import DataValidation
    def op(wb):
        sheet = wb[sheet_name]
        if validation_type == "dropdown":
            formula = f'"{",".join(options)}"'
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            sheet.add_data_validation(dv)
            dv.add(cell_range)
        return True
    return _execute_operation(s3_key, op)

def apply_conditional_formatting(s3_key: str, sheet_name: str, cell_range: str, rule_type: str, condition: str, format_options: dict):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill
    def op(wb):
        sheet = wb[sheet_name]
        fill = PatternFill(start_color=format_options.get("bg_color", "FFFFFF").replace("#", ""), fill_type="solid")
        rule = CellIsRule(operator=rule_type, formula=[condition], fill=fill)
        sheet.conditional_formatting.add(cell_range, rule)
        return True
    return _execute_operation(s3_key, op)

def add_chart(s3_key: str, sheet_name: str, chart_type: str, data_range: str, title: str):
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    def op(wb):
        sheet = wb[sheet_name]
        if chart_type == "bar": chart = BarChart()
        elif chart_type == "line": chart = LineChart()
        elif chart_type == "pie": chart = PieChart()
        else: return False
        
        chart.title = title
        # Very basic ref strategy
        import re
        match = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", data_range)
        if match:
            min_col = openpyxl.utils.column_index_from_string(match.group(1))
            min_row = int(match.group(2))
            max_col = openpyxl.utils.column_index_from_string(match.group(3))
            max_row = int(match.group(4))
            data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
            sheet.add_chart(chart, "E5")
        return True
    return _execute_operation(s3_key, op)

def freeze_panes(s3_key: str, sheet_name: str, row: int, col: int):
    def op(wb):
        sheet = wb[sheet_name]
        cell = f"{get_column_letter(col)}{row}"
        sheet.freeze_panes = cell
        return True
    return _execute_operation(s3_key, op)

def set_row_height(s3_key: str, sheet_name: str, row: int, height: float):
    def op(wb):
        wb[sheet_name].row_dimensions[row].height = height
        return True
    return _execute_operation(s3_key, op)

def set_column_width(s3_key: str, sheet_name: str, column: str, width: float):
    def op(wb):
        wb[sheet_name].column_dimensions[column].width = width
        return True
    return _execute_operation(s3_key, op)
